from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


AZUL = "174F67"
AZUL_CLARO = "DDEBF1"
GRIS = "66727F"
BLANCO = "FFFFFF"
BORDE = Side(style="thin", color="DCE3E8")
FORMATO_MONEDA = '"RD$"#,##0.00;[Red]-"RD$"#,##0.00'


def _valor_excel(valor):
    if isinstance(valor, datetime) and timezone.is_aware(valor):
        return timezone.localtime(valor).replace(tzinfo=None)
    return valor


def _formato_excel(celda, tipo):
    if tipo == "moneda":
        celda.number_format = FORMATO_MONEDA
    elif tipo == "fecha":
        celda.number_format = "dd/mm/yyyy"
    elif tipo == "fecha_hora":
        celda.number_format = "dd/mm/yyyy hh:mm"
    elif tipo == "entero":
        celda.number_format = "0"


def exportar_excel(reporte):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.sheet_view.showGridLines = False
    max_columnas = max(
        [len(seccion.columnas) for seccion in reporte.secciones] + [2]
    )

    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_columnas)
    titulo = hoja.cell(1, 1, reporte.titulo)
    titulo.font = Font(size=18, bold=True, color=BLANCO)
    titulo.fill = PatternFill("solid", fgColor=AZUL)
    titulo.alignment = Alignment(vertical="center")
    hoja.row_dimensions[1].height = 30
    hoja.cell(2, 1, "Edificio")
    hoja.cell(2, 2, reporte.edificio.nombre)
    hoja.cell(3, 1, "Período consultado")
    hoja.cell(3, 2, reporte.periodo_consultado)
    hoja.cell(4, 1, "Fecha de generación")
    hoja.cell(4, 2, timezone.localtime(reporte.generado_en).strftime("%d/%m/%Y %H:%M"))
    for fila in range(2, 5):
        hoja.cell(fila, 1).font = Font(bold=True, color=GRIS)

    fila_actual = 6
    hoja.cell(fila_actual, 1, "Totales conciliados").font = Font(
        bold=True, color=AZUL
    )
    fila_actual += 1
    for total in reporte.totales:
        hoja.cell(fila_actual, 1, total.etiqueta).font = Font(bold=True)
        celda = hoja.cell(fila_actual, 2, _valor_excel(total.valor))
        _formato_excel(celda, total.tipo)
        fila_actual += 1

    fila_actual += 1
    primera_tabla = True
    for seccion in reporte.secciones:
        hoja.merge_cells(
            start_row=fila_actual,
            start_column=1,
            end_row=fila_actual,
            end_column=len(seccion.columnas),
        )
        celda_seccion = hoja.cell(fila_actual, 1, seccion.titulo)
        celda_seccion.font = Font(size=13, bold=True, color=AZUL)
        celda_seccion.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        fila_actual += 1
        fila_encabezado = fila_actual
        for columna, definicion in enumerate(seccion.columnas, start=1):
            celda = hoja.cell(fila_actual, columna, definicion.titulo)
            celda.font = Font(bold=True, color=BLANCO)
            celda.fill = PatternFill("solid", fgColor=AZUL)
            celda.alignment = Alignment(wrap_text=True)
            celda.border = Border(bottom=BORDE)
            letra = get_column_letter(columna)
            hoja.column_dimensions[letra].width = max(
                hoja.column_dimensions[letra].width or 0,
                definicion.ancho,
            )
        fila_actual += 1
        if seccion.filas:
            for fila in seccion.filas:
                for columna, (valor, definicion) in enumerate(
                    zip(fila, seccion.columnas), start=1
                ):
                    celda = hoja.cell(fila_actual, columna, _valor_excel(valor))
                    celda.border = Border(bottom=BORDE)
                    celda.alignment = Alignment(vertical="top", wrap_text=True)
                    _formato_excel(celda, definicion.tipo)
                fila_actual += 1
        else:
            hoja.merge_cells(
                start_row=fila_actual,
                start_column=1,
                end_row=fila_actual,
                end_column=len(seccion.columnas),
            )
            hoja.cell(fila_actual, 1, "Sin registros para los filtros seleccionados")
            hoja.cell(fila_actual, 1).font = Font(italic=True, color=GRIS)
            fila_actual += 1
        if primera_tabla:
            hoja.freeze_panes = f"A{fila_encabezado + 1}"
            primera_tabla = False
        fila_actual += 2

    hoja.print_title_rows = "1:4"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.page_margins.left = 0.3
    hoja.page_margins.right = 0.3
    hoja.page_margins.top = 0.5
    hoja.page_margins.bottom = 0.5
    libro.properties.title = reporte.titulo
    libro.properties.subject = f"{reporte.edificio.nombre} - {reporte.periodo_consultado}"
    libro.properties.creator = "Mi Condominio"

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida.getvalue()


def formatear_valor(valor, tipo):
    if tipo == "moneda":
        return f"RD${Decimal(valor):,.2f}"
    if tipo == "fecha" and isinstance(valor, (date, datetime)):
        return valor.strftime("%d/%m/%Y")
    if tipo == "fecha_hora" and isinstance(valor, datetime):
        return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M")
    if tipo == "entero":
        return f"{int(valor):,}"
    return str(valor)


def exportar_pdf(reporte):
    from weasyprint import HTML

    secciones = []
    for seccion in reporte.secciones:
        filas = [
            [
                formatear_valor(valor, definicion.tipo)
                for valor, definicion in zip(fila, seccion.columnas)
            ]
            for fila in seccion.filas
        ]
        secciones.append(
            {
                "titulo": seccion.titulo,
                "columnas": seccion.columnas,
                "filas": filas,
            }
        )
    totales = [
        {
            "etiqueta": total.etiqueta,
            "valor": formatear_valor(total.valor, total.tipo),
        }
        for total in reporte.totales
    ]
    html = render_to_string(
        "reports/reporte_pdf.html",
        {
            "reporte": reporte,
            "secciones_pdf": secciones,
            "totales_pdf": totales,
        },
    )
    return HTML(string=html).write_pdf()
