from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from accounts.models import UsuarioEdificio
from buildings.models import Apartamento, Edificio
from finance.models import Cargo, Gasto, Pago
from finance.reporting import generar_reporte
from finance.services import (
    confirmar_gasto,
    confirmar_pago,
    generar_cuotas_mensuales,
    registrar_gasto,
    registrar_pago,
    registrar_saldo_inicial,
)
from people.models import Persona, RelacionApartamento


class ReportesExportablesTests(TestCase):
    def setUp(self):
        usuarios = get_user_model()
        self.edificio = Edificio.objects.create(
            nombre="Residencial Reportes",
            direccion="Santiago de los Caballeros",
        )
        self.otro_edificio = Edificio.objects.create(
            nombre="Residencial Confidencial",
            direccion="Santiago",
        )
        self.apartamento = Apartamento.objects.create(
            edificio=self.edificio,
            numero="A-1",
            cuota_mensual=Decimal("1000.00"),
        )
        self.otro_apartamento = Apartamento.objects.create(
            edificio=self.edificio,
            numero="A-2",
            cuota_mensual=Decimal("1500.00"),
        )
        self.apartamento_ajeno = Apartamento.objects.create(
            edificio=self.otro_edificio,
            numero="SECRETO-9",
            cuota_mensual=Decimal("9999.00"),
        )
        self.admin = usuarios.objects.create_user(username="admin-reportes")
        self.residente = usuarios.objects.create_user(username="residente-reportes")
        self.admin_ajeno = usuarios.objects.create_user(username="admin-reporte-ajeno")
        UsuarioEdificio.objects.bulk_create(
            [
                UsuarioEdificio(
                    usuario=self.admin,
                    edificio=self.edificio,
                    rol=UsuarioEdificio.Rol.ADMINISTRADOR,
                ),
                UsuarioEdificio(
                    usuario=self.residente,
                    edificio=self.edificio,
                    rol=UsuarioEdificio.Rol.RESIDENTE,
                ),
                UsuarioEdificio(
                    usuario=self.admin_ajeno,
                    edificio=self.otro_edificio,
                    rol=UsuarioEdificio.Rol.ADMINISTRADOR,
                ),
            ]
        )
        persona = Persona.objects.create(
            usuario=self.residente,
            nombre_completo="Residente Reportes",
            telefono="809-555-0202",
            correo="reportes@example.test",
        )
        RelacionApartamento.objects.create(
            persona=persona,
            apartamento=self.apartamento,
            tipo=RelacionApartamento.Tipo.PROPIETARIO,
            fecha_inicio=timezone.localdate() - timedelta(days=300),
            puede_acceder_portal=True,
            puede_ver_estado_cuenta=True,
        )
        self.periodo, _ = generar_cuotas_mensuales(
            edificio=self.edificio,
            anio=2026,
            mes=9,
            fecha_vencimiento=date(2026, 9, 10),
            usuario=self.admin,
        )
        self.cargo = Cargo.objects.get(
            apartamento=self.apartamento,
            periodo=self.periodo,
        )
        registrar_saldo_inicial(
            edificio=self.edificio,
            importe=Decimal("10000.00"),
            fecha_corte=date(2026, 8, 31),
            concepto="Saldo bancario",
            usuario=self.admin,
        )
        registrar_saldo_inicial(
            edificio=self.edificio,
            apartamento=self.apartamento,
            importe=Decimal("100.00"),
            fecha_corte=date(2026, 8, 31),
            concepto="Balance anterior",
            usuario=self.admin,
        )
        pago = registrar_pago(
            apartamento=self.apartamento,
            importe_total=Decimal("400.00"),
            fecha=date(2026, 9, 5),
            metodo=Pago.Metodo.TRANSFERENCIA,
            referencia="REP-400",
            usuario=self.admin,
        )
        confirmar_pago(
            pago=pago,
            usuario=self.admin,
            aplicaciones=[(self.cargo, Decimal("400.00"))],
        )
        registrar_pago(
            apartamento=self.otro_apartamento,
            importe_total=Decimal("250.00"),
            fecha=date(2026, 9, 8),
            metodo=Pago.Metodo.DEPOSITO,
            referencia="PEND-250",
            usuario=self.admin,
        )
        gasto = registrar_gasto(
            edificio=self.edificio,
            categoria=Gasto.Categoria.MANTENIMIENTO,
            concepto="Bomba de agua",
            importe=Decimal("200.00"),
            fecha=date(2026, 9, 7),
            usuario=self.admin,
        )
        confirmar_gasto(gasto=gasto, usuario=self.admin)
        self.parametros_base = {
            "periodo": str(self.periodo.pk),
            "apartamento": str(self.apartamento.pk),
            "fecha_desde": "2026-09-01",
            "fecha_hasta": "2026-09-30",
        }

    def parametros(self, tipo, formato):
        return {**self.parametros_base, "tipo": tipo, "formato": formato}

    def test_los_siete_reportes_excel_abren_y_conservan_metadatos(self):
        self.client.force_login(self.admin)
        tipos = (
            "estado_cuenta",
            "cuentas_cobrar",
            "cuotas_periodo",
            "gastos_categoria",
            "resumen_mensual",
            "movimientos_caja",
            "pagos_pendientes",
        )
        for tipo in tipos:
            with self.subTest(tipo=tipo):
                response = self.client.get(
                    reverse("finance:reporte_exportar", args=[self.edificio.pk]),
                    self.parametros(tipo, "xlsx"),
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.content[:2], b"PK")
                libro = load_workbook(BytesIO(response.content), data_only=False)
                hoja = libro["Reporte"]
                self.assertEqual(hoja["B2"].value, "Residencial Reportes")
                self.assertEqual(hoja["B3"].value, "09/2026")
                textos = " ".join(
                    str(celda.value)
                    for fila in hoja.iter_rows()
                    for celda in fila
                    if celda.value is not None
                )
                self.assertIn("Fecha de generación", textos)
                self.assertNotIn("Residencial Confidencial", textos)
                self.assertNotIn("SECRETO-9", textos)
                libro.close()

    def test_los_siete_reportes_pdf_abren_correctamente(self):
        self.client.force_login(self.admin)
        tipos = (
            "estado_cuenta",
            "cuentas_cobrar",
            "cuotas_periodo",
            "gastos_categoria",
            "resumen_mensual",
            "movimientos_caja",
            "pagos_pendientes",
        )
        for tipo in tipos:
            with self.subTest(tipo=tipo):
                response = self.client.get(
                    reverse("finance:reporte_exportar", args=[self.edificio.pk]),
                    self.parametros(tipo, "pdf"),
                )
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.content.startswith(b"%PDF-"))
                lector = PdfReader(BytesIO(response.content))
                self.assertGreaterEqual(len(lector.pages), 1)
                texto = " ".join(pagina.extract_text() or "" for pagina in lector.pages)
                self.assertIn("Residencial Reportes", texto)
                self.assertIn("09/2026", texto)
                self.assertNotIn("Residencial Confidencial", texto)
                self.assertNotIn("SECRETO-9", texto)

    def test_estado_de_cuenta_reconcilia_totales(self):
        reporte = generar_reporte(
            tipo="estado_cuenta",
            edificio=self.edificio,
            usuario=self.admin,
            apartamento=self.apartamento,
        )

        totales = {total.etiqueta: total.valor for total in reporte.totales}
        self.assertEqual(totales["Saldo inicial"], Decimal("100.00"))
        self.assertEqual(totales["Cargos"], Decimal("1000.00"))
        self.assertEqual(totales["Pagos aplicados"], Decimal("400.00"))
        self.assertEqual(
            totales["Saldo pendiente conciliado"], Decimal("700.00")
        )

    def test_residente_solo_exporta_estado_de_su_apartamento(self):
        self.client.force_login(self.residente)
        propio = self.client.get(
            reverse("finance:reporte_exportar", args=[self.edificio.pk]),
            self.parametros("estado_cuenta", "xlsx"),
        )
        otro = self.client.get(
            reverse("finance:reporte_exportar", args=[self.edificio.pk]),
            {
                **self.parametros("estado_cuenta", "xlsx"),
                "apartamento": self.otro_apartamento.pk,
            },
        )
        general = self.client.get(
            reverse("finance:reporte_exportar", args=[self.edificio.pk]),
            self.parametros("cuentas_cobrar", "xlsx"),
        )

        self.assertEqual(propio.status_code, 200)
        self.assertEqual(otro.status_code, 400)
        self.assertEqual(general.status_code, 400)

    def test_usuario_de_otro_edificio_no_accede_a_reportes(self):
        self.client.force_login(self.admin_ajeno)

        response = self.client.get(
            reverse("finance:reporte_exportar", args=[self.edificio.pk]),
            self.parametros("cuentas_cobrar", "xlsx"),
        )

        self.assertEqual(response.status_code, 404)
